import sys
import subprocess

# Auto-install dependencies jika script dijalankan langsung tanpa venv
required_modules = ["openpyxl", "playwright"]
missing_modules = []
for mod in required_modules:
    try:
        __import__(mod)
    except ImportError:
        missing_modules.append(mod)

if missing_modules:
    print(f"[Info] Modul belum lengkap ({', '.join(missing_modules)}). Menginstall otomatis...")
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", *missing_modules])
        if "playwright" in missing_modules:
            subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])
        print("[Info] Instalasi dependencies selesai!")
    except Exception as e:
        print(f"[Peringatan] Gagal auto-install modul: {e}")

from playwright.sync_api import sync_playwright
import re
import time
import json
import os
import random
import openpyxl
import zipfile
import xml.etree.ElementTree as ET

CACHE_FILE = "processed_status_ditemukan.json"
REPORT_JSON = "laporan_status_ditemukan.json"
REPORT_EXCEL = "laporan_status_ditemukan.xlsx"

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r") as f:
                return set(json.load(f))
        except Exception:
            return set()
    return set()

def save_cache(cache_set):
    with open(CACHE_FILE, "w") as f:
        json.dump(list(cache_set), f, indent=4)

def load_reports():
    if os.path.exists(REPORT_JSON):
        try:
            with open(REPORT_JSON, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []

def save_report(reports_list):
    try:
        with open(REPORT_JSON, "w", encoding="utf-8") as f:
            json.dump(reports_list, f, indent=4, ensure_ascii=False)
    except Exception as e:
        print(f"  -> Error menyimpan report JSON: {e}")
    try:
        if reports_list:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Laporan"
            headers = list(reports_list[0].keys())
            ws.append(headers)
            for item in reports_list:
                ws.append([item.get(h, "") for h in headers])
            wb.save(REPORT_EXCEL)
    except Exception as e:
        pass

def check_usaha_sudah_diganti(page):
    """
    Mengecek apakah terdapat peringatan bahwa usaha ini pernah ditautkan tetapi sudah diganti dengan yang lain,
    sehingga saat ini tidak ditemukan lagi pada keluarga tersebut.
    """
    try:
        res = page.evaluate('''() => {
            const bodyText = (document.body.innerText || '').toLowerCase();
            const hasKeyword1 = bodyText.includes('pernah ditautkan') || bodyText.includes('ditautkan pada usaha');
            const hasKeyword2 = bodyText.includes('sudah diganti') || bodyText.includes('tidak ditemukan lagi pada keluarga') || bodyText.includes('tidak ditemukan lagi');
            
            if (hasKeyword1 && hasKeyword2) {
                const allElements = Array.from(document.querySelectorAll('div, p, span, [role="alert"], .alert'));
                const alertEl = allElements.find(el => {
                    const t = (el.innerText || el.textContent || '').trim().toLowerCase();
                    return t.includes('pernah ditautkan') && (t.includes('sudah diganti') || t.includes('tidak ditemukan lagi'));
                });
                const fullText = alertEl ? alertEl.innerText.trim() : "Usaha ini pernah ditautkan pada usaha keluarga, namun usaha yang ditautkan sudah diganti sehingga saat ini tidak ditemukan lagi pada keluarga tersebut.";
                return { detected: true, message: fullText };
            }
            return { detected: false, message: "" };
        }''')
        return res.get("detected", False), res.get("message", "")
    except Exception:
        return False, ""

def check_is_bot_or_blocked(page):
    """Mengecek apakah halaman saat ini menunjukkan pesan terdeteksi bot, WAF, atau error SSO."""
    try:
        if page.locator("text=/mendeteksi koneksi anda sebagai bot/i").count() > 0:
            return True
        if page.locator("text=/HaloSIS/i").count() > 0:
            return True
        if page.locator("text=/Lanjutkan dengan SSO/i").count() > 0:
            return True
        if page.locator("text=/Access Denied/i").count() > 0:
            return True
        url = page.url.lower()
        if "sso" in url and ("error" in url or "login" in url or "block" in url):
            return True
    except:
        pass
    return False

def resolve_bot_detection(page, target_link):
    """
    Penanganan terdeteksi bot (WAF/SSO):
    Tunggu lebih tenang agar IP rate-limit ter-reset, tidak reload agresif.
    """
    if not check_is_bot_or_blocked(page):
        return True

    print("="*60)
    print(" [WAF / BOT DETECTED] Halaman terdeteksi bot atau terganggu sesi SSO.")
    print(" Memulai prosedur pemulihan otomatis...")
    
    max_retries = 2
    for attempt in range(1, max_retries + 1):
        print(f" -> [Percobaan {attempt}/{max_retries}] Menunggu 10 detik agar rate-limit ter-reset...")
        time.sleep(10)
        
        print(" -> Refresh/reload halaman...")
        try:
            page.reload(wait_until="networkidle", timeout=15000)
        except Exception:
            try:
                page.goto(target_link, wait_until="networkidle", timeout=15000)
            except Exception:
                pass
        time.sleep(4)
        
        # Cari dan klik tombol 'Lanjutkan dengan SSO' jika ada
        try:
            sso_btn = page.locator("button, a, div").filter(
                has_text=re.compile(r"Lanjutkan dengan SSO|Lanjutkan SSO|Login.*SSO|Masuk.*SSO|Lanjutkan", re.IGNORECASE)
            )
            if sso_btn.count() > 0 and sso_btn.first.is_visible():
                print(" -> Mengklik tombol 'Lanjutkan dengan SSO'...")
                sso_btn.first.click()
                time.sleep(5)
                try:
                    page.wait_for_load_state("networkidle", timeout=10000)
                except Exception:
                    pass
        except Exception as e:
            print(f" -> Cek tombol SSO info: {e}")

        if not check_is_bot_or_blocked(page):
            print(" -> [BERHASIL] Halaman terbebas dari deteksi bot! Melanjutkan bot...")
            print("="*60)
            return True

    print("="*60)
    print(" TERDETEKSI SEBAGAI BOT OLEH SERVER (WAF BPS / HALOSIS)!")
    print(" Coba klik tombol 'Kembali' di browser atau ganti koneksi internet (Tethering HP).")
    print(" Tekan ENTER di terminal ini jika halaman sudah kembali normal.")
    print("="*60)
    input("Tekan ENTER untuk melanjutkan bot...")
    try:
        page.goto(target_link, wait_until="networkidle", timeout=15000)
    except Exception:
        pass
    return True

def get_hidden_rows(excel_path):
    """
    Mendeteksi baris yang disembunyikan (hidden) atau ter-filter (AutoFilter) di Excel
    secara super cepat langsung dari struktur XML arsip XLSX (< 1 detik).
    """
    hidden_rows = set()
    try:
        with zipfile.ZipFile(excel_path, "r") as z:
            # 1. Parse workbook.xml & rels untuk mencari worksheet XML yang tepat
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
            rels_xml = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            
            rel_map = {}
            for rel in rels_xml:
                r_id = rel.attrib.get("Id")
                target = rel.attrib.get("Target")
                if r_id and target:
                    if not target.startswith("xl/"):
                        target = "xl/" + target.lstrip("/")
                    rel_map[r_id] = target
            
            target_xml_path = None
            ns = {
                "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            }
            
            sheets = wb_xml.findall(".//main:sheet", ns)
            for s in sheets:
                s_name = s.attrib.get("name", "")
                r_id = s.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                # Ambil sheet bernama 'data' atau sheet pertama sebagai target
                if s_name.lower() == "data" or target_xml_path is None:
                    target_xml_path = rel_map.get(r_id)
                    if s_name.lower() == "data":
                        break
            
            if target_xml_path and target_xml_path in z.namelist():
                sheet_tree = ET.fromstring(z.read(target_xml_path))
                for r in sheet_tree.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row"):
                    # Baris tersembunyi jika hidden="1" atau tinggi baris ht="0"
                    if r.attrib.get("hidden") == "1" or r.attrib.get("ht") == "0":
                        hidden_rows.add(int(r.attrib["r"]))
    except Exception as e:
        print(f"[Info] Deteksi filter XML info: {e}")
    return hidden_rows

def read_excel_data(excel_path):
    """Membaca file Excel secara streaming cepat (read_only=True) dengan dukungan filter Excel (AutoFilter / Hidden Rows)."""
    print(f"Membaca file: {os.path.basename(excel_path)} ...")
    
    # Deteksi baris yang disembunyikan/difilter di Excel
    hidden_rows = get_hidden_rows(excel_path)
    if hidden_rows:
        print(f"-> Mendeteksi filter Excel aktif: {len(hidden_rows)} baris tersembunyi akan otomatis dilewati.")
    
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active

    header_row = 1
    link_col_idx = None
    kab_col_idx = None
    nama_col_idx = None
    status_col_idx = None

    rows_data = []
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx <= 10 and link_col_idx is None:
            # Deteksi header
            row_str = [str(c).strip().lower() if c is not None else '' for c in row]
            for c_idx, val in enumerate(row_str):
                if 'link_fasih' in val or ('link' in val and 'fasih' in val) or val == 'link':
                    link_col_idx = c_idx
                elif 'kab' in val or 'kabupaten' in val:
                    kab_col_idx = c_idx
                elif 'namabangunanusaha' in val or 'namausaha' in val or 'nama bangunan' in val:
                    if nama_col_idx is None:
                        nama_col_idx = c_idx
                elif 'statuskeberadaan' in val or 'status' in val:
                    status_col_idx = c_idx
            
            if link_col_idx is not None:
                header_row = row_idx
                continue

        if link_col_idx is None:
            continue

        if row_idx <= header_row:
            continue

        # Lewati baris yang disembunyikan / difilter di Excel
        if row_idx in hidden_rows:
            continue

        raw_link = row[link_col_idx] if len(row) > link_col_idx else None
        if not raw_link or str(raw_link).strip() == '' or str(raw_link).strip().lower() == 'none':
            continue

        link = str(raw_link).strip()
        # Pastikan URL memiliki akhiran /edit
        if not link.endswith("/edit"):
            link = link.rstrip("/") + "/edit"

        kab = str(row[kab_col_idx]).strip() if (kab_col_idx is not None and len(row) > kab_col_idx and row[kab_col_idx] is not None) else "UNKNOWN"
        nama = str(row[nama_col_idx]).strip() if (nama_col_idx is not None and len(row) > nama_col_idx and row[nama_col_idx] is not None) else ""
        
        rows_data.append({
            "link": link,
            "kab": kab,
            "nama": nama
        })

    wb.close()
    return rows_data

def select_kabupaten(rows_data):
    """Menampilkan pilihan filter Kabupaten/Kota atau seluruh data."""
    kab_counts = {}
    for item in rows_data:
        k = item["kab"]
        kab_counts[k] = kab_counts.get(k, 0) + 1

    sorted_kabs = sorted(kab_counts.keys())
    if len(sorted_kabs) <= 1:
        return rows_data

    print("\n" + "="*65)
    print(" PILIHAN FILTER KABUPATEN / KOTA:")
    print("="*65)
    print(f" [0] SEMUA KABUPATEN / KOTA ({len(rows_data)} total data)")
    for idx, k in enumerate(sorted_kabs, 1):
        cnt = kab_counts[k]
        print(f" [{idx}] Kab/Kota: {k} ({cnt} data)")
    print("="*65)
    
    choice = input("Pilih nomor (atau ketik langsung kode Kab, default [0] Semua): ").strip()
    if not choice or choice == "0":
        print("-> Memproses seluruh Kabupaten/Kota.")
        return rows_data

    # Cek jika input adalah kode kab langsung (misal: 3301)
    if choice in kab_counts:
        selected_kab = choice
        filtered = [item for item in rows_data if item["kab"] == selected_kab]
        print(f"-> Filter aktif: Kab {selected_kab} ({len(filtered)} data).")
        return filtered

    # Cek jika input adalah index nomor urut
    try:
        idx_choice = int(choice)
        if 1 <= idx_choice <= len(sorted_kabs):
            selected_kab = sorted_kabs[idx_choice - 1]
            filtered = [item for item in rows_data if item["kab"] == selected_kab]
            print(f"-> Filter aktif: Kab {selected_kab} ({len(filtered)} data).")
            return filtered
    except ValueError:
        pass

    print("-> Pilihan tidak valid, memproses SEMUA data.")
    return rows_data

def process_single_assignment(page, item):
    """
    Melakukan alur perubahan status menjadi '1. Ditemukan' dan Submit Paksa:
    1. Buka URL /edit
    2. Verifikasi halaman /edit (jika dialihkan, skip)
    3. Klik menu 'SE2026 - P'
    4. Pilih opsi '1. Ditemukan' pada pertanyaan Keberadaan Bangunan Lainnya/ Usaha
    5. Klik tombol 'Kirim'
    6. Klik tombol titik tiga (split dropdown) pada modal
    7. Klik menuitem 'Submit Paksa'
    8. Konfirmasi jika ada dialog konfirmasi lanjutan
    """
    link = item["link"]
    
    # 1. Navigasi ke URL /edit
    try:
        page.goto(link, wait_until="networkidle", timeout=30000)
    except Exception as e:
        if "interrupted by another navigation" in str(e):
            time.sleep(2)
            page.goto(link, wait_until="networkidle", timeout=30000)
        else:
            raise e

    # Tangani WAF/SSO jika muncul
    resolve_bot_detection(page, link)

    # 2. Verifikasi berada di halaman /edit
    current_url = page.url
    if "/edit" not in current_url:
        print(f"  -> [SKIP] Browser dialihkan ke {current_url}. Bukan wilayah tugas Anda / form terkunci.")
        return "SKIP_NOT_AUTHORIZED", "Bukan wilayah tugas Anda / form terkunci"

    # Cek apakah ada peringatan usaha sudah diganti
    is_diganti, msg_diganti = check_usaha_sudah_diganti(page)
    if is_diganti:
        return "SKIP_USAHA_SUDAH_DIGANTI", msg_diganti

    # 3. Klik menu 'SE2026 - P'
    print("  -> Mencari dan mengklik menu 'SE2026 - P'...")
    menu_clicked = False

    def is_se2026_rendered():
        try:
            return page.evaluate('''() => {
                const text = (document.body.innerText || '').toLowerCase();
                return text.includes("keberadaan bangunan") || text.includes("1. ditemukan") || text.includes("tidak ditemukan");
            }''')
        except Exception:
            return False

    # Jika form sudah menampilkan pertanyaan SE2026 - P, tidak perlu klik ulang
    if is_se2026_rendered():
        print("     [OK] Form 'SE2026 - P' sudah aktif/terbuka.")
        menu_clicked = True
    else:
        pattern = re.compile(r"SE2026\s*[-–—]?\s*P", re.IGNORECASE)
        exact_pattern = re.compile(r"^\s*SE2026\s*[-–—]?\s*P\s*$", re.IGNORECASE)

        # Strategi 1: Cari dengan selector Playwright presisi (tanpa generic div!)
        candidates = [
            page.get_by_role("tab", name=pattern),
            page.get_by_role("link", name=pattern),
            page.get_by_role("button", name=pattern),
            page.get_by_text(exact_pattern),
            page.locator("a, button, [role='tab'], [role='menuitem'], [role='treeitem'], li").filter(has_text=pattern),
            page.locator("[title*='SE2026 - P'], [title*='SE2026-P'], [title*='SE2026']"),
            page.get_by_text(pattern)
        ]

        for cand in candidates:
            try:
                cnt = cand.count()
                if cnt > 0:
                    target_el = None
                    for c_i in range(cnt):
                        el = cand.nth(c_i)
                        if el.is_visible():
                            target_el = el
                            break
                    
                    if target_el:
                        target_el.scroll_into_view_if_needed(timeout=2000)
                        box = target_el.bounding_box()
                        if box:
                            page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
                        else:
                            target_el.click(timeout=3000)
                        
                        time.sleep(1.5)
                        if is_se2026_rendered():
                            menu_clicked = True
                            print("     [OK] Berhasil berpindah ke menu 'SE2026 - P' via Playwright.")
                            break
            except Exception:
                pass

        # Strategi 2: JS Deep Click (Cari leaf node terdalam, scroll, dan dispatch pointer+mouse events)
        if not menu_clicked:
            js_clicked = page.evaluate('''() => {
                const all = Array.from(document.querySelectorAll('*'));
                const matches = all.filter(el => {
                    const txt = (el.innerText || el.textContent || '').trim();
                    return /SE2026\s*[-–—]?\s*P/i.test(txt);
                });
                
                if (matches.length === 0) return { success: false, reason: 'not_found' };

                // Urutkan dari elemen dengan teks terpendek (elemen terdalam / leaf)
                matches.sort((a, b) => {
                    const aLen = (a.innerText || a.textContent || '').trim().length;
                    const bLen = (b.innerText || b.textContent || '').trim().length;
                    return aLen - bLen;
                });

                const leaf = matches[0];
                const clickable = leaf.closest('a, button, [role="tab"], [role="treeitem"], [role="menuitem"], li, [tabindex]') || leaf;
                
                clickable.scrollIntoView({ behavior: 'instant', block: 'center' });
                
                const rect = clickable.getBoundingClientRect();
                const x = rect.left + rect.width / 2;
                const y = rect.top + rect.height / 2;
                const opts = { bubbles: true, cancelable: true, view: window, clientX: x, clientY: y };
                
                clickable.dispatchEvent(new PointerEvent('pointerdown', opts));
                clickable.dispatchEvent(new MouseEvent('mousedown', opts));
                clickable.dispatchEvent(new PointerEvent('pointerup', opts));
                clickable.dispatchEvent(new MouseEvent('mouseup', opts));
                clickable.click();
                
                return {
                    success: true,
                    tag: clickable.tagName,
                    text: (clickable.innerText || clickable.textContent || '').trim().substring(0, 30)
                };
            }''')

            if js_clicked and js_clicked.get("success"):
                time.sleep(2)
                if is_se2026_rendered():
                    menu_clicked = True
                    print(f"     [OK] Berhasil berpindah ke menu 'SE2026 - P' via JS deep click ({js_clicked.get('tag')}).")

        # Strategi 3: Tunggu render tambahan jika aplikasi butuh waktu memuat form
        if not menu_clicked:
            for _ in range(3):
                time.sleep(1)
                if is_se2026_rendered():
                    menu_clicked = True
                    break

        if not menu_clicked:
            print("  -> [Warning] Tampilan belum terkonfirmasi berubah ke form 'SE2026 - P'. Mencoba mencari opsi pertanyaan...")

    time.sleep(1) # Beri jeda form stabil

    # Cek apakah ada peringatan bahwa usaha sudah diganti pada form SE2026 - P
    is_diganti, msg_diganti = check_usaha_sudah_diganti(page)
    if is_diganti:
        return "SKIP_USAHA_SUDAH_DIGANTI", msg_diganti

    # 4. Pilih opsi '1. Ditemukan'
    print("  -> Memilih opsi '1. Ditemukan' pada Keberadaan Bangunan Lainnya/ Usaha...")
    option_selected = False

    # Tunggu beberapa saat jika form sedang proses rendering
    for _ in range(8):
        if page.evaluate('''() => {
            const body = (document.body.innerText || '').toLowerCase();
            return body.includes('1. ditemukan') || body.includes('ditemukan');
        }'''):
            break
        time.sleep(0.5)

    # Coba pendekatan JS cerdas
    select_result = page.evaluate('''() => {
        function triggerChange(el) {
            el.dispatchEvent(new Event('change', { bubbles: true }));
            el.dispatchEvent(new Event('input', { bubbles: true }));
        }

        const allElements = Array.from(document.querySelectorAll('*'));
        
        // 1. Cari radio input dengan value="1"
        const radios = Array.from(document.querySelectorAll('input[type="radio"]'));
        const radioOne = radios.find(r => r.value === '1' || r.value === '1. Ditemukan');
        if (radioOne) {
            radioOne.scrollIntoView({ behavior: 'instant', block: 'center' });
            radioOne.click();
            radioOne.checked = true;
            triggerChange(radioOne);
            const parentLabel = radioOne.closest('label');
            if (parentLabel) parentLabel.click();
            return { success: true, method: 'radio_input_val_1' };
        }

        // 2. Cari elemen teks yang mengandung "1. Ditemukan"
        const foundTexts = allElements.filter(el => {
            const t = (el.innerText || el.textContent || '').trim();
            return /^1\.\s*ditemukan/i.test(t);
        });

        if (foundTexts.length > 0) {
            foundTexts.sort((a, b) => (a.innerText || '').length - (b.innerText || '').length);
            const targetEl = foundTexts[0];
            const clickable = targetEl.closest('label, div[role="radio"], button, tr, div') || targetEl;
            
            clickable.scrollIntoView({ behavior: 'instant', block: 'center' });
            clickable.click();
            
            const innerRadio = clickable.querySelector('input[type="radio"]');
            if (innerRadio) {
                innerRadio.checked = true;
                triggerChange(innerRadio);
            }
            return { success: true, method: 'label_or_text_click' };
        }

        return { success: false };
    }''')

    if select_result and select_result.get("success"):
        option_selected = True
        print(f"     [OK] Berhasil memilih '1. Ditemukan' ({select_result.get('method')})")
    else:
        # Fallback Playwright
        playwright_selectors = [
            page.get_by_role("radio", name=re.compile(r"1\.\s*Ditemukan", re.IGNORECASE)),
            page.get_by_label(re.compile(r"1\.\s*Ditemukan", re.IGNORECASE)),
            page.locator("label").filter(has_text=re.compile(r"1\.\s*Ditemukan", re.IGNORECASE)),
            page.locator("div, span, button").filter(has_text=re.compile(r"^1\.\s*Ditemukan", re.IGNORECASE))
        ]
        for sel in playwright_selectors:
            try:
                if sel.count() > 0 and sel.first.is_visible():
                    sel.first.scroll_into_view_if_needed(timeout=2000)
                    box = sel.first.bounding_box()
                    if box:
                        page.mouse.click(box["x"] + 10, box["y"] + box["height"] / 2)
                    else:
                        sel.first.click(timeout=3000)
                    option_selected = True
                    print("     [OK] Berhasil memilih '1. Ditemukan' via Playwright fallback.")
                    break
            except Exception:
                pass

    if not option_selected:
        diag_info = page.evaluate('''() => {
            const titles = Array.from(document.querySelectorAll('h1, h2, h3, h4, [role="tab"], button'))
                .map(e => (e.innerText || '').trim())
                .filter(t => t.length > 0 && t.length < 40)
                .slice(0, 15);
            return { currentUrl: window.location.href, titles: titles };
        }''')
        print(f"  -> [Diagnostic] Elemen di halaman saat ini: {diag_info}")
        raise Exception("Gagal memilih '1. Ditemukan' pada form.")

    time.sleep(1)

    # 5. Klik tombol 'Kirim'
    print("  -> Mengklik tombol 'Kirim'...")
    kirim_clicked = False
    for attempt in range(3):
        try:
            kirim_btn = page.get_by_role("button", name=re.compile(r"^Kirim$", re.IGNORECASE)).first
            if kirim_btn.is_visible():
                kirim_btn.click(timeout=5000)
                kirim_clicked = True
                break
        except Exception:
            pass

        # Fallback JS untuk klik tombol Kirim
        found_kirim_js = page.evaluate('''() => {
            const btns = Array.from(document.querySelectorAll('button'));
            const target = btns.find(b => {
                const t = (b.innerText || b.textContent || '').trim().toLowerCase();
                return (t === 'kirim' || t.includes('kirim')) && b.offsetParent !== null;
            });
            if (target) {
                target.click();
                return true;
            }
            return false;
        }''')
        if found_kirim_js:
            kirim_clicked = True
            break
        time.sleep(1)

    if not kirim_clicked:
        raise Exception("Gagal mengklik tombol 'Kirim' pada form.")

    time.sleep(1.5) # Tunggu modal muncul

    # 6. Klik tombol titik tiga (split button dropdown) pada modal
    print("  -> Mengklik tombol menu titik tiga (split dropdown) pada modal...")
    
    # Tunggu modal muncul dan selesai animasi
    try:
        page.locator("[role='dialog'], .modal, div[data-state='open']").first.wait_for(state="visible", timeout=5000)
    except Exception:
        pass
    time.sleep(1)

    dots_clicked = False

    def is_menu_open():
        try:
            return page.evaluate('''() => {
                const items = Array.from(document.querySelectorAll('[role="menuitem"], [role="menu"] *, div, button, span'));
                return items.some(el => {
                    const t = (el.innerText || el.textContent || '').trim().toLowerCase();
                    return t.includes('submit paksa') && el.offsetParent !== null;
                });
            }''')
        except Exception:
            return False

    split_selectors = [
        page.locator("button:has(svg path[d*='M12 12m-1'])"),
        page.locator("button.tw\\:rounded-r-lg, button[class*='rounded-r-lg']"),
        page.locator("[role='dialog'] button:has(svg path[d*='M12 12m-1'])"),
        page.locator("[role='dialog'] button.tw\\:rounded-r-lg"),
        page.locator("[role='dialog'] button:has(svg)")
    ]

    for attempt in range(5):
        if is_menu_open():
            dots_clicked = True
            print("     [OK] Menu 'Submit Paksa' sudah muncul.")
            break

        # 1. Coba klik via Playwright native mouse click pada bounding box tombol
        clicked_via_pw = False
        for loc in split_selectors:
            try:
                cnt = loc.count()
                if cnt > 0:
                    for i in range(cnt):
                        el = loc.nth(i)
                        if el.is_visible():
                            box = el.bounding_box()
                            if box and box["width"] > 0:
                                page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
                                clicked_via_pw = True
                                break
                            else:
                                el.click(timeout=2000, force=True)
                                clicked_via_pw = True
                                break
                    if clicked_via_pw:
                        break
            except Exception:
                pass

        time.sleep(1)
        if is_menu_open():
            dots_clicked = True
            print("     [OK] Menu titik tiga berhasil dibuka via Playwright mouse click.")
            break

        # 2. Coba klik via JS dengan event simulation lengkap (pointerdown -> mousedown -> click)
        page.evaluate('''() => {
            const btns = Array.from(document.querySelectorAll('button'));
            
            let target = btns.find(b => {
                const c = b.className || '';
                return c.includes('rounded-r-lg') && b.offsetParent !== null;
            });

            if (!target) {
                target = btns.find(b => {
                    const svgs = b.querySelectorAll('svg');
                    for (const s of svgs) {
                        const html = s.innerHTML || '';
                        if (html.includes('M12 12m-1') || html.includes('M12 5m-1') || html.includes('M12 19m-1')) {
                            return true;
                        }
                    }
                    return false;
                });
            }

            if (!target) {
                const modal = document.querySelector('[role="dialog"], .modal, div[data-state="open"]');
                if (modal) {
                    const modalBtns = Array.from(modal.querySelectorAll('button'));
                    target = modalBtns.find(b => {
                        const rect = b.getBoundingClientRect();
                        return rect.width > 10 && rect.width < 50 && rect.height > 20;
                    });
                }
            }

            if (target) {
                target.scrollIntoView({ behavior: 'instant', block: 'center' });
                const rect = target.getBoundingClientRect();
                const x = rect.left + rect.width / 2;
                const y = rect.top + rect.height / 2;
                const opts = { bubbles: true, cancelable: true, view: window, clientX: x, clientY: y, pointerId: 1, pointerType: 'mouse', isPrimary: true };
                
                target.dispatchEvent(new PointerEvent('pointerdown', opts));
                target.dispatchEvent(new MouseEvent('mousedown', opts));
                target.dispatchEvent(new PointerEvent('pointerup', opts));
                target.dispatchEvent(new MouseEvent('mouseup', opts));
                target.click();
            }
        }''')

        time.sleep(1)
        if is_menu_open():
            dots_clicked = True
            print("     [OK] Menu titik tiga berhasil dibuka via JS event simulation.")
            break

    if not is_menu_open():
        # Cek apakah ada peringatan usaha sudah diganti yang baru terdeteksi
        is_diganti, msg_diganti = check_usaha_sudah_diganti(page)
        if is_diganti:
            return "SKIP_USAHA_SUDAH_DIGANTI", msg_diganti

        diag_btns = page.evaluate('''() => {
            return Array.from(document.querySelectorAll('button')).map(b => ({
                text: (b.innerText || '').trim(),
                className: b.className || '',
                visible: b.offsetParent !== null
            })).filter(b => b.visible);
        }''')
        print(f"  -> [Diagnostic] Daftar tombol terlihat di modal: {diag_btns}")
        raise Exception("Tombol titik tiga (split dropdown) gagal memunculkan menu 'Submit Paksa'.")

    # 7. Klik menu item 'Submit Paksa'
    print("  -> Mengklik 'Submit Paksa'...")
    submit_paksa_clicked = False

    paksa_locators = [
        page.locator("[role='menuitem']").filter(has_text=re.compile(r"Submit Paksa", re.IGNORECASE)),
        page.get_by_text(re.compile(r"^Submit Paksa$", re.IGNORECASE)),
        page.locator("div, button, a, span").filter(has_text=re.compile(r"Submit Paksa", re.IGNORECASE))
    ]

    for loc in paksa_locators:
        try:
            cnt = loc.count()
            if cnt > 0:
                for i in range(cnt):
                    el = loc.nth(i)
                    if el.is_visible():
                        box = el.bounding_box()
                        if box:
                            page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
                        else:
                            el.click(timeout=2000, force=True)
                        submit_paksa_clicked = True
                        print("     [OK] Berhasil mengklik 'Submit Paksa' via Playwright.")
                        break
            if submit_paksa_clicked:
                break
        except Exception:
            pass

    if not submit_paksa_clicked:
        found_paksa_js = page.evaluate('''() => {
            const items = Array.from(document.querySelectorAll('[role="menuitem"], div, button, span'));
            const target = items.find(el => {
                const t = (el.innerText || el.textContent || '').trim().toLowerCase();
                return t.includes('submit paksa') && el.offsetParent !== null;
            });
            if (target) {
                target.scrollIntoView({ behavior: 'instant', block: 'center' });
                const rect = target.getBoundingClientRect();
                const x = rect.left + rect.width / 2;
                const y = rect.top + rect.height / 2;
                const opts = { bubbles: true, cancelable: true, view: window, clientX: x, clientY: y, pointerId: 1, pointerType: 'mouse', isPrimary: true };
                
                target.dispatchEvent(new PointerEvent('pointerdown', opts));
                target.dispatchEvent(new MouseEvent('mousedown', opts));
                target.dispatchEvent(new PointerEvent('pointerup', opts));
                target.dispatchEvent(new MouseEvent('mouseup', opts));
                target.click();
                return true;
            }
            return false;
        }''')
        if found_paksa_js:
            submit_paksa_clicked = True
            print("     [OK] Berhasil mengklik 'Submit Paksa' via JS simulation.")

    if not submit_paksa_clicked:
        raise Exception("Gagal mengklik menu 'Submit Paksa'.")

    time.sleep(1.5)

    # 8. Cek dan tangani dialog konfirmasi lanjutan jika ada
    # (Misal: 'Apakah Anda yakin ingin submit paksa?', tombol 'Ya' / 'Konfirmasi' / 'Submit')
    try:
        page.evaluate('''() => {
            const btns = Array.from(document.querySelectorAll('button'));
            const confirmBtn = btns.find(b => {
                const t = (b.innerText || b.textContent || '').trim().toLowerCase();
                return (t === 'ya' || t === 'konfirmasi' || t === 'submit' || t === 'setuju' || t.includes('ya, submit')) && b.offsetParent !== null;
            });
            if (confirmBtn) {
                confirmBtn.click();
            }
        }''')
    except Exception:
        pass

    time.sleep(2.5) # Tunggu pengiriman selesai ke server Fasih
    return "SUCCESS", "Status diubah ke Ditemukan & Submit Paksa sukses"

def main():
    print("="*65)
    print(" BOT UBAH STATUS ANOMALI KE 'DITEMUKAN' & SUBMIT PAKSA")
    print(" Fasih-SM BPS (SE2026)")
    print("="*65)

    data_dir = os.path.join(os.getcwd(), "data")
    if not os.path.exists(data_dir):
        os.makedirs(data_dir)
        print(f"Folder 'data' telah dibuat di {data_dir}")
        print("Silakan masukkan file Excel anomali ke folder 'data' lalu jalankan ulang.")
        return

    excel_files = [f for f in os.listdir(data_dir) if f.endswith(".xlsx") and not f.startswith("~$")]
    if not excel_files:
        print("TIDAK ADA FILE EXCEL DITEMUKAN di folder 'data'!")
        print(f"Silakan letakkan file Excel Anda di: {data_dir}")
        return

    # Prioritaskan file '12. BKU Ditautkan tapi Status Ganda.xlsx' jika ada
    target_file = None
    for f in excel_files:
        if "12" in f and "ganda" in f.lower():
            target_file = f
            break
    
    if target_file is None:
        if len(excel_files) == 1:
            target_file = excel_files[0]
        else:
            print("\nFile Excel yang tersedia:")
            for idx, f in enumerate(excel_files, 1):
                print(f" [{idx}] {f}")
            choice = input(f"Pilih file [1-{len(excel_files)}] (default 1): ").strip()
            try:
                c_idx = int(choice) - 1
                target_file = excel_files[c_idx] if 0 <= c_idx < len(excel_files) else excel_files[0]
            except ValueError:
                target_file = excel_files[0]

    excel_path = os.path.join(data_dir, target_file)
    rows_data = read_excel_data(excel_path)
    if not rows_data:
        print("Tidak ada baris link yang valid di file tersebut.")
        return

    print(f"-> Total {len(rows_data)} baris data berhasil dibaca dari {target_file}.")

    # Pilihan filter kabupaten/kota
    active_data = select_kabupaten(rows_data)

    processed_cache = load_cache()
    pending_items = [item for item in active_data if item["link"] not in processed_cache]
    already_done = len(active_data) - len(pending_items)
    pct = (already_done / len(active_data) * 100) if active_data else 0.0

    print("\n" + "="*65)
    print(" REKAPAN STATUS PENGERJAAN:")
    print("="*65)
    print(f" - File Digunakan     : {target_file}")
    print(f" - Total Link Target  : {len(active_data)}")
    print(f" - Sudah Selesai      : {already_done} ({pct:.1f}%)")
    print(f" - Sisa Diproses      : {len(pending_items)}")
    print("="*65)

    if not pending_items:
        print("Semua data pada filter ini sudah berhasil diproses!")
        return

    print("\nMembuka browser Playwright...")
    with sync_playwright() as p:
        user_data_dir = os.path.join(os.getcwd(), "chrome_profile_anomali")
        context = p.chromium.launch_persistent_context(
            user_data_dir=user_data_dir,
            headless=False,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--start-maximized"
            ],
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            no_viewport=True
        )

        page = context.pages[0] if context.pages else context.new_page()
        page.add_init_script("delete navigator.__proto__.webdriver;")

        page.goto("https://fasih-sm.bps.go.id/")
        print("Silakan pastikan Anda telah login ke Fasih-SM (SSO BPS).")
        print("="*65)
        input("Tekan ENTER di terminal ini jika sudah siap memulai proses...")

        reports_data = load_reports()
        report_link_map = {r.get("link"): idx for idx, r in enumerate(reports_data)}

        def append_or_update_report(item_data):
            link = item_data.get("link")
            if link in report_link_map:
                reports_data[report_link_map[link]] = item_data
            else:
                report_link_map[link] = len(reports_data)
                reports_data.append(item_data)
            save_report(reports_data)

        total_target = len(active_data)
        success_count = 0
        skip_count = 0
        skip_diganti_count = 0
        fail_count = 0

        for idx, item in enumerate(active_data, 1):
            link = item["link"]
            kab = item["kab"]
            nama = item["nama"]

            if link in processed_cache:
                print(f"[{idx}/{total_target}] [Kab {kab}] SKIP (Sudah ada di cache): {nama}")
                continue

            print(f"[{idx}/{total_target}] [Kab {kab}] Memproses: {nama} ({link})")

            try:
                status, msg = process_single_assignment(page, item)
                
                if status == "SUCCESS":
                    processed_cache.add(link)
                    save_cache(processed_cache)
                    success_count += 1
                    append_or_update_report({
                        "waktu": time.strftime("%Y-%m-%d %H:%M:%S"),
                        "kab": kab,
                        "nama": nama,
                        "link": link,
                        "status": "BERHASIL",
                        "keterangan": "Status diubah ke Ditemukan & Submit Paksa sukses",
                        "detail_pesan": msg
                    })
                    print(f"  -> [BERHASIL] Status diubah ke Ditemukan & Submit Paksa sukses!")
                
                elif status == "SKIP_USAHA_SUDAH_DIGANTI":
                    processed_cache.add(link)
                    save_cache(processed_cache)
                    skip_diganti_count += 1
                    append_or_update_report({
                        "waktu": time.strftime("%Y-%m-%d %H:%M:%S"),
                        "kab": kab,
                        "nama": nama,
                        "link": link,
                        "status": "SKIP - USAHA SUDAH DIGANTI",
                        "keterangan": "Usaha sudah diganti dengan yang lain, jadi assignment ini tidak ditemukan",
                        "detail_pesan": msg
                    })
                    print(f"  -> [SKIP] Usaha sudah diganti dengan yang lain (assignment tidak ditemukan).")
                    print(f"     Pesan: {msg}")

                elif status == "SKIP_NOT_AUTHORIZED":
                    processed_cache.add(link)
                    save_cache(processed_cache)
                    skip_count += 1
                    append_or_update_report({
                        "waktu": time.strftime("%Y-%m-%d %H:%M:%S"),
                        "kab": kab,
                        "nama": nama,
                        "link": link,
                        "status": "SKIP - BUKAN OTORISASI",
                        "keterangan": "Bukan wilayah tugas Anda / form terkunci",
                        "detail_pesan": msg
                    })
                    print(f"  -> [SKIP] Ditandai di cache (Bukan otorisasi Anda).")

            except Exception as e:
                fail_count += 1
                print(f"  -> [GAGAL] Error saat memproses: {e}")
                if check_is_bot_or_blocked(page):
                    resolve_bot_detection(page, link)
                print("     Lanjut ke baris berikutnya...")

            # Jeda acak antar link (6-11 detik)
            delay = random.uniform(6, 11)
            time.sleep(delay)

            # Istirahat 30-45 detik setiap kelipatan 15 link agar tidak terkena rate limit WAF
            if idx % 15 == 0 and idx < total_target:
                rest_time = random.uniform(30, 45)
                print("="*60)
                print(f" [COOLING DOWN] Telah memproses {idx} data. Jeda sejenak selama {int(rest_time)} detik...")
                print("="*60)
                time.sleep(rest_time)

        print("\n" + "="*65)
        print(" PROSES OTOMATISASI SELESAI!")
        print(f" - Sukses Diproses           : {success_count}")
        print(f" - Usaha Sudah Diganti (Skip): {skip_diganti_count}")
        print(f" - Bukan Otorisasi (Skip)    : {skip_count}")
        print(f" - Gagal/Error               : {fail_count}")
        print(f" - File Laporan Rekap        : {REPORT_EXCEL} & {REPORT_JSON}")
        print("="*65)
        context.close()

if __name__ == "__main__":
    main()
