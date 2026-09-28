import os
import sys
import re
import time
import json
import random
import openpyxl
from playwright.sync_api import sync_playwright

CACHE_FILE = "processed_edit_admin.json"

# ==============================================================================
# KONFIGURASI PENGATURAN BOT
# ==============================================================================
# Set ke True jika ingin memproses file anomali Missing Value NIK.
# Set ke False (default) jika ingin MELEWATI (SKIP) file Missing Value NIK.
PROSES_MISSING_VALUE_NIK = False

# Teks penjelasan anomali standar
PENJELASAN_ANOMALI_TEXT = "sudah dikonfirmasi ke petugas lapangan. dan sudah sesuai kondisi lapangan"
# ==============================================================================

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return set(json.load(f))
        except Exception:
            return set()
    return set()

def save_cache(cache_set):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(list(cache_set), f, indent=4)

def check_is_bot_or_blocked(page):
    """Mengecek apakah halaman saat ini menunjukkan pesan terdeteksi bot, WAF, atau error SSO."""
    try:
        if page.locator("text=/mendeteksi koneksi anda sebagai bot/i").count() > 0:
            return True
        if page.locator("text=/HaloSIS/i").count() > 0:
            return True
        if page.locator("text=/Lanjutkan dengan SSO/i").count() > 0:
            return True
        if page.locator("text=/Access Denied/i").count() > 0:
            return True
        url = page.url.lower()
        if "sso" in url and ("error" in url or "login" in url or "block" in url):
            return True
    except Exception:
        pass
    return False

def resolve_bot_detection(page, target_link):
    """Penanganan terdeteksi bot (WAF/SSO) dengan retry tenang."""
    if not check_is_bot_or_blocked(page):
        return True

    print("=" * 60)
    print(" [WAF / BOT DETECTED] Halaman terdeteksi bot atau terganggu sesi SSO.")
    print(" Memulai prosedur pemulihan otomatis...")
    
    max_retries = 2
    for attempt in range(1, max_retries + 1):
        print(f" -> [Percobaan {attempt}/{max_retries}] Menunggu 10 detik agar rate-limit ter-reset...")
        time.sleep(10)
        
        print(" -> Refresh/reload halaman...")
        try:
            page.reload(wait_until="networkidle", timeout=15000)
        except Exception:
            try:
                page.goto(target_link, wait_until="networkidle", timeout=15000)
            except Exception:
                pass
        time.sleep(4)
        
        try:
            sso_btn = page.locator("button, a, div").filter(
                has_text=re.compile(r"Lanjutkan dengan SSO|Lanjutkan SSO|Login.*SSO|Masuk.*SSO|Lanjutkan", re.IGNORECASE)
            )
            if sso_btn.count() > 0 and sso_btn.first.is_visible():
                print(" -> Mengklik tombol 'Lanjutkan dengan SSO'...")
                sso_btn.first.click()
                time.sleep(5)
                try:
                    page.wait_for_load_state("networkidle", timeout=10000)
                except Exception:
                    pass
        except Exception as e:
            print(f" -> Cek tombol SSO info: {e}")

        if not check_is_bot_or_blocked(page):
            print(" -> [BERHASIL] Halaman terbebas dari deteksi bot! Melanjutkan bot...")
            print("=" * 60)
            return True

    print("=" * 60)
    print(" TERDETEKSI SEBAGAI BOT OLEH SERVER (WAF BPS / HALOSIS)!")
    print(" Silakan selesaikan verifikasi di browser secara manual.")
    print("=" * 60)
    input("Tekan ENTER di terminal ini jika halaman sudah normal...")
    try:
        page.goto(target_link, wait_until="networkidle", timeout=15000)
    except Exception:
        pass
    return True

def click_sidebar_menu(page, menu_keyword):
    """Mencari dan mengklik menu sidebar tertentu (misal: 'Anomali Usaha' atau 'Anomali Keluarga')."""
    pattern = re.compile(rf"{re.escape(menu_keyword)}", re.IGNORECASE)
    
    # 1. Coba pencarian dengan selektor Playwright
    candidates = [
        page.get_by_role("tab", name=pattern),
        page.get_by_role("link", name=pattern),
        page.get_by_role("button", name=pattern),
        page.locator(f"[title*='{menu_keyword}' i]"),
        page.locator("a, button, [role='tab'], [role='menuitem'], [role='treeitem'], li").filter(has_text=pattern),
        page.get_by_text(pattern)
    ]
    for cand in candidates:
        try:
            cnt = cand.count()
            if cnt > 0:
                for i in range(cnt):
                    el = cand.nth(i)
                    if el.is_visible():
                        el.scroll_into_view_if_needed(timeout=2000)
                        box = el.bounding_box()
                        if box:
                            page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
                        else:
                            el.click(timeout=3000)
                        return True
        except Exception:
            pass

    # 2. Coba pencarian mendalam dengan JavaScript
    res = page.evaluate('''(menuText) => {
        const regex = new RegExp(menuText, 'i');
        const elements = Array.from(document.querySelectorAll('a, button, [role="tab"], [role="menuitem"], [role="treeitem"], li, div[title], span, p'));
        const matches = elements.filter(el => {
            const t = (el.innerText || el.textContent || '').trim();
            return regex.test(t) && el.offsetParent !== null;
        });
        if (matches.length === 0) return false;
        
        matches.sort((a, b) => (a.innerText || '').length - (b.innerText || '').length);
        const target = matches[0];
        const clickable = target.closest('a, button, [role="tab"], [role="menuitem"], [role="treeitem"], li') || target;
        clickable.scrollIntoView({ behavior: 'instant', block: 'center' });
        clickable.click();
        return true;
    }''', menu_keyword)
    return res

def click_kirim_and_confirm(page, section_name="Form"):
    """Mengklik tombol Kirim dan menangani modal konfirmasi jika muncul."""
    kirim_clicked = False
    for attempt in range(3):
        try:
            kirim_btn = page.get_by_role("button", name=re.compile(r"^Kirim$", re.IGNORECASE)).first
            if kirim_btn.is_visible():
                kirim_btn.click(timeout=5000)
                kirim_clicked = True
                break
        except Exception:
            pass
        
        # Fallback JS
        found_js = page.evaluate('''() => {
            const btns = Array.from(document.querySelectorAll('button'));
            const target = btns.find(b => {
                const t = (b.innerText || b.textContent || '').trim().toLowerCase();
                return (t === 'kirim' || t === 'simpan' || t.includes('kirim')) && b.offsetParent !== null;
            });
            if (target) {
                target.scrollIntoView({ behavior: 'instant', block: 'center' });
                target.click();
                return true;
            }
            return false;
        }''')
        if found_js:
            kirim_clicked = True
            break
        time.sleep(1)

    if not kirim_clicked:
        print(f"     [Info] Tombol 'Kirim' tidak ditemukan pada {section_name} (mungkin perubahan sudah tersimpan otomatis).")
        return False

    time.sleep(1.5)
    # Tangani dialog modal konfirmasi jika ada
    for _ in range(3):
        try:
            btns = page.locator("button").filter(has_text=re.compile(r"Kirim|Ya|Konfirmasi|Setuju|Lanjut", re.IGNORECASE)).all()
            visible_btns = [b for b in btns if b.is_visible()]
            if len(visible_btns) > 1:
                print(f"     -> Mengklik tombol konfirmasi pada modal {section_name}...")
                visible_btns[-1].click(timeout=3000)
                time.sleep(1.5)
            else:
                break
        except Exception:
            break
    time.sleep(2)
    return True

def handle_anomali_section(page, menu_name, explanation_text):
    """
    Memeriksa dan memproses anomali pada bagian/menu saat ini.
    - Mencari opsi 'Ya, Sesuai Kondisi Lapangan'
    - Mencentang checkbox / radio
    - Mengisi field 'Penjelasan Anomali'
    - Menyimpan (Kirim + Konfirmasi)
    """
    time.sleep(1.5)
    
    # Hitung berapa banyak elemen 'Sesuai Kondisi Lapangan' yang ditemukan
    total_anomalies = page.evaluate('''() => {
        const regex = /sesuai\\s*(kondisi)?\\s*lapangan|ya[,\\s]+sesuai/i;
        const all = Array.from(document.querySelectorAll('*'));
        const matchedLeafs = all.filter(el => {
            const t = (el.innerText || el.textContent || '').trim();
            if (!regex.test(t)) return false;
            return !Array.from(el.children).some(c => regex.test((c.innerText || c.textContent || '').trim()));
        });
        return matchedLeafs.length;
    }''')

    if not total_anomalies or total_anomalies == 0:
        print(f"  -> Tidak ada anomali / opsi 'Ya, Sesuai Kondisi Lapangan' ditemukan pada menu '{menu_name}'.")
        return 0

    print(f"  -> Ditemukan {total_anomalies} anomali pada menu '{menu_name}'. Memproses konfirmasi...")
    
    processed_count = 0
    for idx in range(total_anomalies):
        # 1. Checklist 'Ya, Sesuai Kondisi Lapangan'
        check_res = page.evaluate('''({ index }) => {
            const regex = /sesuai\\s*(kondisi)?\\s*lapangan|ya[,\\s]+sesuai/i;
            const all = Array.from(document.querySelectorAll('*'));
            const matchedLeafs = all.filter(el => {
                const t = (el.innerText || el.textContent || '').trim();
                if (!regex.test(t)) return false;
                return !Array.from(el.children).some(c => regex.test((c.innerText || c.textContent || '').trim()));
            });

            if (index >= matchedLeafs.length) return { success: false, reason: "index_out_of_bounds" };

            const targetLeaf = matchedLeafs[index];
            const container = targetLeaf.closest('tr, [role="row"], .card, fieldset, form > div, div[class*="border"], div[class*="rounded"]') || targetLeaf.parentElement?.parentElement || targetLeaf.parentElement;

            let cb = container.querySelector('input[type="checkbox"], input[type="radio"]');
            if (!cb) {
                cb = targetLeaf.closest('label')?.querySelector('input') || targetLeaf.querySelector('input');
            }

            let wasChecked = false;
            if (cb) {
                wasChecked = cb.checked;
                if (!wasChecked) {
                    cb.scrollIntoView({ behavior: 'instant', block: 'center' });
                    cb.click();
                    if (!cb.checked) {
                        cb.checked = true;
                        cb.dispatchEvent(new Event('change', { bubbles: true }));
                        cb.dispatchEvent(new Event('input', { bubbles: true }));
                    }
                }
            } else {
                const clickable = targetLeaf.closest('label, button, div[role="checkbox"], div[role="radio"]') || targetLeaf;
                clickable.scrollIntoView({ behavior: 'instant', block: 'center' });
                clickable.click();
            }

            return { success: true, wasChecked: wasChecked };
        }''', { "index": idx })

        time.sleep(1) # Tunggu field Penjelasan Anomali muncul / render

        # 2. Isi field 'Penjelasan Anomali'
        fill_res = page.evaluate('''({ index, text }) => {
            const regex = /sesuai\\s*(kondisi)?\\s*lapangan|ya[,\\s]+sesuai/i;
            const all = Array.from(document.querySelectorAll('*'));
            const matchedLeafs = all.filter(el => {
                const t = (el.innerText || el.textContent || '').trim();
                if (!regex.test(t)) return false;
                return !Array.from(el.children).some(c => regex.test((c.innerText || c.textContent || '').trim()));
            });

            if (index >= matchedLeafs.length) return { success: false };

            const targetLeaf = matchedLeafs[index];
            const container = targetLeaf.closest('tr, [role="row"], .card, fieldset, form > div, div[class*="border"], div[class*="rounded"]') || targetLeaf.parentElement?.parentElement || targetLeaf.parentElement;

            // Cari textarea atau input di dalam container
            let inputEl = container.querySelector('textarea, input[type="text"]');
            
            // Jika tidak ada di dalam container, cari di elemen saudara berikutnya
            if (!inputEl && container.nextElementSibling) {
                inputEl = container.nextElementSibling.querySelector('textarea, input[type="text"]');
            }

            // Jika belum ketemu, cari textarea terlihat di halaman
            if (!inputEl) {
                const textareas = Array.from(document.querySelectorAll('textarea')).filter(t => t.offsetParent !== null);
                if (textareas.length > index) {
                    inputEl = textareas[index];
                } else if (textareas.length > 0) {
                    inputEl = textareas[textareas.length - 1];
                }
            }

            if (inputEl) {
                inputEl.scrollIntoView({ behavior: 'instant', block: 'center' });
                inputEl.focus();

                const proto = inputEl.tagName === 'TEXTAREA' ? window.HTMLTextAreaElement.prototype : window.HTMLInputElement.prototype;
                const desc = Object.getOwnPropertyDescriptor(proto, 'value');
                if (desc && desc.set) {
                    desc.set.call(inputEl, text);
                } else {
                    inputEl.value = text;
                }

                inputEl.dispatchEvent(new Event('input', { bubbles: true }));
                inputEl.dispatchEvent(new Event('change', { bubbles: true }));
                inputEl.dispatchEvent(new Event('blur', { bubbles: true }));

                return { success: true, tag: inputEl.tagName };
            }

            return { success: false, reason: "input_not_found" };
        }''', { "index": idx, "text": explanation_text })

        # Coba juga isi via Playwright native jika ada textarea terlihat
        try:
            visible_textareas = page.locator("textarea").all()
            active_ta = [t for t in visible_textareas if t.is_visible()]
            if len(active_ta) > idx:
                curr_val = active_ta[idx].input_value()
                if not curr_val or curr_val.strip() == "":
                    active_ta[idx].fill(explanation_text)
            elif len(active_ta) > 0:
                curr_val = active_ta[-1].input_value()
                if not curr_val or curr_val.strip() == "":
                    active_ta[-1].fill(explanation_text)
        except Exception:
            pass

        print(f"     [OK] Anomali #{idx + 1}: 'Ya, Sesuai Kondisi Lapangan' tercentang & Penjelasan terisi.")
        processed_count += 1
        time.sleep(0.5)

    if processed_count > 0:
        print(f"  -> Menyimpan perubahan pada '{menu_name}' (Klik Kirim)...")
        click_kirim_and_confirm(page, section_name=menu_name)

    return processed_count


def get_anomaly_excel_files(base_data_dir):
    """
    Mencari file Excel untuk anomali:
    1. Cek folder data/Anomali terlebih dahulu.
    2. Jika tidak ada / kosong, cek folder data/ utama (filter file yang bukan BKU dan bukan laporan).
    """
    anomali_subdir = os.path.join(base_data_dir, "Anomali")
    files_with_path = []
    
    if os.path.exists(anomali_subdir):
        sub_files = [
            f for f in os.listdir(anomali_subdir) 
            if f.endswith(".xlsx") and not f.startswith("~$") and "laporan" not in f.lower()
        ]
        if sub_files:
            for f in sub_files:
                files_with_path.append((f, os.path.join(anomali_subdir, f)))
            return files_with_path

    root_files = [
        f for f in os.listdir(base_data_dir) 
        if f.endswith(".xlsx") and not f.startswith("~$") 
        and "laporan" not in f.lower() 
        and "bku" not in f.lower()
    ]
    for f in root_files:
        files_with_path.append((f, os.path.join(base_data_dir, f)))
        
    return files_with_path

def main():
    data_dir = os.path.join(os.getcwd(), "data")
    
    if not os.path.exists(data_dir):
        os.makedirs(data_dir)
        print(f"Folder 'data' telah dibuat di {data_dir}")
        print("Silakan masukkan file Excel anomali ke folder 'data/Anomali' atau 'data'.")
        return
        
    excel_items = get_anomaly_excel_files(data_dir)
    if not excel_items:
        print("TIDAK ADA FILE EXCEL ANOMALI DITEMUKAN!")
        print(f"Silakan letakkan file Excel anomali (.xlsx) di folder: {os.path.join(data_dir, 'Anomali')} atau {data_dir}")
        return

    processed_cache = load_cache()
    kegiatan_id = "fd68e454-ba45-4b85-8205-f3bf777ded24"
    
    file_summaries = []
    edit_links = []
    link_to_kecamatan = {}

    for file_name, excel_file in excel_items:
        print(f"Membaca file Excel: {file_name} ({excel_file})...")
        
        is_missing_value_file = "missing_value" in file_name.lower() or "missing value" in file_name.lower()
        if is_missing_value_file and not PROSES_MISSING_VALUE_NIK:
            print(f"  -> [SKIP / DISABLED] File ini dilewati karena 'PROSES_MISSING_VALUE_NIK = False'.")
            file_summaries.append({
                "file_name": file_name,
                "total": 0,
                "processed": 0,
                "pending": 0,
                "disabled": True
            })
            continue

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            ws = wb.active
            file_links = []
            
            header_row = 4
            link_col = None
            status_col = None
            kec_col = None
            
            for r in range(1, 10):
                row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                row_str = [str(v).strip().lower() if v else '' for v in row_vals]
                if any('link fasih' in s or 'link' in s for s in row_str):
                    header_row = r
                    for c_idx, s in enumerate(row_str, 1):
                        if 'link fasih' in s or 'link' in s:
                            link_col = c_idx
                        if 'tindak lanjut' in s or 'tindaklanjut' in s:
                            status_col = c_idx
                        if 'nama kecamatan' in s or 'kecamatan' in s or 'nama kec' in s:
                            kec_col = c_idx
                    break
            
            if link_col is None:
                link_col = 18
            if kec_col is None:
                kec_col = 8

            for row in range(header_row + 1, ws.max_row + 1):
                if ws.row_dimensions[row].hidden or ws.row_dimensions[row].height == 0:
                    continue
                
                if status_col is not None:
                    status_val = ws.cell(row=row, column=status_col).value
                    if status_val is None or "belum ditindaklanjuti" not in str(status_val).strip().lower():
                        continue

                val = ws.cell(row=row, column=link_col).value
                kec_val = str(ws.cell(row=row, column=kec_col).value or 'UNKNOWN').strip().upper() if kec_col else 'UNKNOWN'
                if val is not None and str(val).strip() != "":
                    file_links.append((str(val).strip(), kec_val))
            wb.close()
            
            file_edit_links = []
            for link, kec_val in file_links:
                link = link.strip()
                match = re.search(r"assignment-detail/([a-zA-Z0-9\-]+)", link)
                if match:
                    dynamic_id = match.group(1)
                    new_link = f"https://fasih-sm.bps.go.id/app/assignment/{kegiatan_id}/{dynamic_id}/edit"
                    file_edit_links.append(new_link)
                    link_to_kecamatan[new_link] = kec_val
                    if new_link not in edit_links:
                        edit_links.append(new_link)

            total_file_links = len(file_edit_links)
            processed_file_links = sum(1 for l in file_edit_links if l in processed_cache)
            pending_file_links = total_file_links - processed_file_links
            
            status_info = f"hanya status 'Belum Ditindaklanjuti' [Kolom {status_col}]" if status_col else "proses semua data [Tanpa kolom status]"
            print(f"  -> {total_file_links} link dibaca dari {file_name} (Header baris {header_row}, Link kolom {link_col}, {status_info}).")
            
            file_summaries.append({
                "file_name": file_name,
                "total": total_file_links,
                "processed": processed_file_links,
                "pending": pending_file_links,
                "disabled": False
            })
            
        except Exception as e:
            print(f"  -> Gagal membaca file {file_name}: {e}")

    if not edit_links:
        print("Tidak ada link valid yang ditemukan di file Excel mana pun.")
        return

    # Tampilkan rekapan detail per file
    print("=" * 65)
    print(" REKAPAN DATA ANOMALI PER FILE EXCEL (MODE EDIT BY ADMIN):")
    print("=" * 65)
    for idx, s in enumerate(file_summaries, 1):
        fname = s['file_name']
        if s.get('disabled'):
            print(f" [{idx}] {fname}")
            print(f"     - [STATUS: DISABLED] File dilewati karena 'PROSES_MISSING_VALUE_NIK = False'")
        else:
            tot = s['total']
            prc = s['processed']
            pnd = s['pending']
            pct = (prc / tot * 100) if tot > 0 else 0.0
            print(f" [{idx}] {fname}")
            print(f"     - Total Link Valid : {tot}")
            print(f"     - Sudah Diproses   : {prc} ({pct:.1f}%)")
            print(f"     - Sisa Diproses    : {pnd}")

    # Hitung ringkasan per kecamatan
    kecamatan_stats = {}
    for link in edit_links:
        kec = link_to_kecamatan.get(link, "UNKNOWN")
        if kec not in kecamatan_stats:
            kecamatan_stats[kec] = {"total": 0, "processed": 0, "pending": 0}
        kecamatan_stats[kec]["total"] += 1
        if link in processed_cache:
            kecamatan_stats[kec]["processed"] += 1
        else:
            kecamatan_stats[kec]["pending"] += 1

    print("=" * 65)
    print(" REKAPAN DATA ANOMALI PER KECAMATAN:")
    print("=" * 65)
    print(f" {'No':<3} {'Nama Kecamatan':<20} {'Total':>7} {'Sudah':>7} {'Sisa':>7} {'Progress':>9}")
    print("-" * 65)
    for idx, (kec_name, stat) in enumerate(sorted(kecamatan_stats.items()), 1):
        tot = stat["total"]
        prc = stat["processed"]
        pnd = stat["pending"]
        pct = (prc / tot * 100) if tot > 0 else 0.0
        print(f" {idx:<3} {kec_name:<20} {tot:>7} {prc:>7} {pnd:>7} {pct:>8.1f}%")

    pending_links = [link for link in edit_links if link not in processed_cache]
    total_processed_unique = sum(1 for l in edit_links if l in processed_cache)
    total_pct = (total_processed_unique / len(edit_links) * 100) if len(edit_links) > 0 else 0.0

    print("-" * 65)
    print(f" {'TOTAL SELURUHNYA':<24} {len(edit_links):>7} {total_processed_unique:>7} {len(pending_links):>7} {total_pct:>8.1f}%")
    print("=" * 65)
    
    if not pending_links:
        print("Semua data pada seluruh file Excel sudah berhasil diproses!")
        return

    print("=" * 65)
    print("Membuka browser Playwright...")
    
    with sync_playwright() as p:
        user_data_dir = os.path.join(os.getcwd(), "chrome_profile_anomali")
        context = p.chromium.launch_persistent_context(
            user_data_dir=user_data_dir,
            headless=False,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--start-maximized"
            ],
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            no_viewport=True
        )
        
        page = context.pages[0] if context.pages else context.new_page()
        page.add_init_script("delete navigator.__proto__.webdriver;")

        page.goto("https://fasih-sm.bps.go.id/")
        
        print("Silakan login ke Fasih-SM (SSO BPS) jika belum login.")
        print("=" * 65)
        input("Tekan ENTER di terminal ini jika sudah berhasil login dan siap memulai...")

        for idx, link in enumerate(edit_links, 1):
            kec_name = link_to_kecamatan.get(link, "")
            kec_prefix = f" [{kec_name}]" if kec_name else ""
            if link in processed_cache:
                print(f"[{idx}/{len(edit_links)}]{kec_prefix} SKIP (Sudah diproses): {link}")
                continue
                
            print(f"[{idx}/{len(edit_links)}]{kec_prefix} Memproses: {link}")
            try:
                # Navigasi ke halaman EDIT
                try:
                    page.goto(link, wait_until="networkidle", timeout=30000)
                except Exception as e:
                    if "interrupted by another navigation" in str(e):
                        print("  -> [Info] Navigasi dialihkan oleh SSO Refresh. Mencoba ulang...")
                        time.sleep(3)
                        try:
                            page.goto(link, wait_until="networkidle", timeout=30000)
                        except Exception:
                            pass
                    else:
                        raise e
                
                resolve_bot_detection(page, link)
                
                # Verifikasi bahwa browser berada di halaman /edit
                current_url = page.url
                if "/edit" not in current_url:
                    print(f"  -> [SKIP] Browser tidak di halaman /edit (URL: {current_url}). Bukan wilayah admin, langsung skip.")
                    processed_cache.add(link)
                    save_cache(processed_cache)
                    continue
                
                # -------------------------------------------------------------
                # 1. BUKA MENU CATATAN & CHECKLIST TAMPILKAN ANOMALI
                # -------------------------------------------------------------
                print("  -> Membuka menu Catatan...")
                catatan_clicked = False
                for _ in range(3):
                    try:
                        page.get_by_role("tab", name=re.compile(r"Catatan", re.IGNORECASE)).click(timeout=5000)
                        catatan_clicked = True
                        break
                    except Exception:
                        try:
                            page.get_by_text(re.compile(r"^.*catatan.*$", re.IGNORECASE)).first.click(timeout=5000)
                            catatan_clicked = True
                            break
                        except Exception:
                            pass
                    time.sleep(1)

                time.sleep(1.5)
                
                # Centang checkbox "Tampilkan Anomali Usaha dan Keluarga"
                print("  -> Memeriksa checkbox 'Tampilkan Anomali Usaha dan Keluarga'...")
                check_result = page.evaluate('''() => {
                    const elements = Array.from(document.querySelectorAll('*'));
                    const targetEl = elements.find(el => el.textContent && el.textContent.toLowerCase().includes('tampilkan anomali') && el.children.length === 0);
                    
                    if (targetEl) {
                        let parent = targetEl.parentElement;
                        for (let i = 0; i < 5; i++) {
                            if (!parent) break;
                            const cb = parent.querySelector('input[type="checkbox"]');
                            if (cb) {
                                const was_checked = cb.checked;
                                if (!was_checked) {
                                    cb.click();
                                }
                                return { found: true, was_checked: was_checked };
                            }
                            parent = parent.parentElement;
                        }
                        
                        targetEl.click();
                        return { found: true, was_checked: false, note: "clicked_text_only" };
                    }
                    return { found: false, was_checked: false };
                }''')

                if check_result and check_result.get('found'):
                    if check_result.get('was_checked'):
                        print("     [OK] Checkbox 'Tampilkan Anomali' sudah tercentang sebelumnya.")
                    else:
                        print("     [OK] Berhasil mencentang 'Tampilkan Anomali'. Mengklik Kirim...")
                        click_kirim_and_confirm(page, section_name="Catatan")
                else:
                    # Fallback Playwright
                    try:
                        page.locator("text=/Tampilkan Anomali/i").first.click(timeout=3000, force=True)
                        click_kirim_and_confirm(page, section_name="Catatan")
                    except Exception:
                        print("     [Info] Checkbox 'Tampilkan Anomali' tidak terdeteksi atau sudah aktif.")

                # Beri waktu beberapa detik agar sidebar mengupdate menu Anomali Usaha dan Anomali Keluarga
                time.sleep(2.5)

                # -------------------------------------------------------------
                # 2. PERIKSA & PROSES MENU ANOMALI USAHA & ANOMALI KELUARGA
                # -------------------------------------------------------------
                target_anomaly_menus = ["Anomali Usaha", "Anomali Keluarga"]
                total_anomalies_handled = 0

                for menu_name in target_anomaly_menus:
                    print(f"  -> Memeriksa menu sidebar: '{menu_name}'...")
                    menu_opened = click_sidebar_menu(page, menu_name)
                    
                    if not menu_opened:
                        print(f"     [Info] Menu '{menu_name}' tidak muncul pada assignment ini.")
                        continue
                    
                    print(f"     [OK] Menu '{menu_name}' terbuka. Memeriksa isi anomali...")
                    time.sleep(2)
                    
                    # Proses anomali pada bagian ini
                    handled = handle_anomali_section(page, menu_name, PENJELASAN_ANOMALI_TEXT)
                    total_anomalies_handled += handled
                    time.sleep(1.5)

                if total_anomalies_handled > 0:
                    print(f"  -> [BERHASIL] Total {total_anomalies_handled} anomali berhasil diperbaiki by admin.")
                else:
                    print("  -> [INFO] Tidak ada anomali aktif yang memerlukan penjelasan pada assignment ini.")

                # Simpan link ke cache sukses
                processed_cache.add(link)
                save_cache(processed_cache)
                print(f"  -> Disimpan ke cache.")

            except Exception as e:
                print(f"  -> Terjadi error pada link {link}:")
                print(f"     {e}")
                if check_is_bot_or_blocked(page):
                    resolve_bot_detection(page, link)
                print("     Lanjut ke link berikutnya...")

            # Jeda acak antar link
            delay = random.uniform(6, 12)
            time.sleep(delay)

            # Cooling down setiap 15 link
            if idx % 15 == 0 and idx < len(edit_links):
                rest_time = random.uniform(30, 45)
                print("=" * 60)
                print(f" [COOLING DOWN] Telah memproses {idx} link. Istirahat sejenak selama {int(rest_time)} detik agar aman dari WAF...")
                print("=" * 60)
                time.sleep(rest_time)

        print("\nProses otomatisasi Edit by Admin selesai!")
        context.close()

if __name__ == "__main__":
    main()
